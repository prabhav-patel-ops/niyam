"""Time-bounded extraction in disposable processes, never the API event loop."""
import io
import multiprocessing
import os
import threading
import zipfile

SLOTS=threading.BoundedSemaphore(2)


def _extract(content,ext,pipe):
    try:
        if os.name!='nt':
            import resource
            resource.setrlimit(resource.RLIMIT_AS,(512*1024*1024,512*1024*1024))
            resource.setrlimit(resource.RLIMIT_CPU,(10,10))
        if ext=='.pdf':
            from pypdf import PdfReader
            pdf=PdfReader(io.BytesIO(content))
            if pdf.is_encrypted or len(pdf.pages)>40:
                raise ValueError('Use an unlocked PDF of at most 40 pages.')
            parts=[];size=0
            for i,page in enumerate(pdf.pages):
                part=f'[Page {i+1}] '+(page.extract_text() or '')
                size+=len(part)
                if size>60000:raise ValueError('Attach only the relevant pages (60,000 characters maximum).')
                parts.append(part)
            text='\n'.join(parts)
        elif ext=='.docx':
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if len(archive.infolist())>1000 or sum(x.file_size for x in archive.infolist())>30*1024*1024:
                    raise ValueError('The expanded Word document is too large.')
            from docx import Document
            doc=Document(io.BytesIO(content))
            text='\n'.join([p.text for p in doc.paragraphs]+[' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows])
        else:
            text=content.decode('utf-8-sig')
        if len(text.strip())<30:
            raise ValueError('No readable strategy text was found. Use a text PDF or paste the rules; scanned-image OCR is not enabled.')
        if len(text)>60000:
            raise ValueError('Attach only the relevant pages (60,000 characters maximum).')
        pipe.send((True,text))
    except ValueError as error:
        pipe.send((False,str(error)[:300]))
    except BaseException:
        pipe.send((False,'This document could not be read. Try a text PDF, DOCX or plain text.'))
    finally:
        pipe.close()


def extract(content,ext):
    if ext not in {'.pdf','.docx','.txt','.md'}:
        raise ValueError('Please attach PDF, DOCX, TXT or Markdown.')
    if not SLOTS.acquire(blocking=False):
        raise ValueError('Document reading is busy. Please try again shortly.')
    context=multiprocessing.get_context('spawn')
    parent,child=context.Pipe(duplex=False)
    process=context.Process(target=_extract,args=(content,ext,child),daemon=True)
    try:
        process.start();child.close()
        if not parent.poll(20):
            raise ValueError('This document took too long to read. Attach fewer pages or paste the rules.')
        try:ok,value=parent.recv()
        except EOFError:
            raise ValueError('This document exceeded the safe reading limits. Try fewer pages.')
        if not ok:raise ValueError(value)
        return value
    finally:
        if process.pid is not None:
            if process.is_alive():process.terminate()
            process.join(timeout=2)
        parent.close();child.close();SLOTS.release()
