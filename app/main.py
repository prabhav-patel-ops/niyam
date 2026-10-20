from contextlib import asynccontextmanager
from datetime import date
from collections import defaultdict,deque
from pathlib import Path
import hashlib
import io
import json
import os
import sqlite3
import threading
import time
import zipfile
from dotenv import load_dotenv

ROOT=Path(__file__).resolve().parents[1]
load_dotenv(ROOT/'.env')

from fastapi import FastAPI,Depends,HTTPException,Request,UploadFile,File
from fastapi.responses import JSONResponse,FileResponse
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.cors import CORSMiddleware
from fastapi.middleware.trustedhost import TrustedHostMiddleware
from fastapi.exceptions import RequestValidationError
from fastapi.security import HTTPBearer,HTTPAuthorizationCredentials
from pydantic import BaseModel,EmailStr,Field,ConfigDict
from . import store,signals,jobs,feed
from .models import Strategy,RunRequest,template
from .data import catalog,load_dataset
from .rules import describe
from .assistant import connection_status,interpret,REFERENCES


@asynccontextmanager
async def lifespan(app):
    store.init()
    stop=threading.Event()
    worker=threading.Thread(target=signals.loop,args=(stop,),daemon=True)
    worker.start()
    yield
    stop.set()
    worker.join(timeout=2)


app=FastAPI(title='Niyam API',version='1.0.0',lifespan=lifespan,docs_url='/api/docs')
origins=[x.strip() for x in os.getenv('NIYAM_ALLOWED_ORIGINS','http://localhost:8000,http://127.0.0.1:8000,https://whood-yop.github.io').split(',') if x.strip()]
app.add_middleware(CORSMiddleware,allow_origins=origins,allow_credentials=False,allow_methods=['GET','POST','PUT','DELETE','OPTIONS'],allow_headers=['Authorization','Content-Type'])
app.add_middleware(TrustedHostMiddleware,allowed_hosts=os.getenv('NIYAM_ALLOWED_HOSTS','localhost,127.0.0.1,testserver').split(','))
auth=HTTPBearer(auto_error=False)
rate=defaultdict(deque)
rate_lock=threading.Lock()


def limit(key,max_count=30,seconds=60):
    now=time.monotonic()
    with rate_lock:
        if len(rate)>10000:
            for old in list(rate):
                if not rate[old] or rate[old][-1]<now-3600:
                    del rate[old]
        q=rate[key]
        while q and q[0]<now-seconds:
            q.popleft()
        if len(q)>=max_count:
            raise HTTPException(429,'Please wait a minute before trying again.')
        q.append(now)


@app.middleware('http')
async def safety_headers(request,call_next):
    size=request.headers.get('content-length','0')
    try:
        if int(size)>9*1024*1024:
            return JSONResponse({'detail':'The file is too large. Maximum size: 8 MB.'},status_code=413)
    except ValueError:
        return JSONResponse({'detail':'Invalid request size.'},status_code=400)
    response=await call_next(request)
    response.headers['X-Content-Type-Options']='nosniff'
    response.headers['Referrer-Policy']='no-referrer'
    response.headers['X-Frame-Options']='DENY'
    if request.url.path.startswith('/api/'):
        response.headers['Cache-Control']='no-store'
    return response


@app.exception_handler(RequestValidationError)
async def friendly_validation(request,error):
    errors=[{'field':'.'.join(str(x) for x in e['loc'][1:]),'message':e['msg']} for e in error.errors()]
    return JSONResponse({'detail':'Please check the highlighted fields.','fields':errors},status_code=422)


def current_user(credentials:HTTPAuthorizationCredentials|None=Depends(auth)):
    user=store.user_for(credentials.credentials) if credentials else None
    if not user:
        raise HTTPException(401,'Sign in to save and run your strategies.')
    return dict(user)


class Signup(BaseModel):
    name:str=Field(min_length=1,max_length=60)
    email:EmailStr
    password:str=Field(min_length=12,max_length=256)


class Login(BaseModel):
    email:EmailStr
    password:str=Field(min_length=1,max_length=256)


class Recovery(BaseModel):
    email:EmailStr
    recovery_code:str=Field(min_length=20,max_length=100)
    password:str=Field(min_length=12,max_length=256)


@app.post('/api/auth/signup')
def signup(body:Signup,request:Request):
    limit('signup:'+request.client.host,5,600)
    uid=store.new_id();recovery=store.new_id()+store.new_id()[:12]
    email=str(body.email).lower()
    try:
        with store.db() as c:
            c.execute('INSERT INTO users VALUES(?,?,?,?,?,?)',(uid,email,body.name.strip(),store.password_hash(body.password),time.time(),hashlib.sha256(recovery.encode()).hexdigest()))
    except sqlite3.IntegrityError:
        raise HTTPException(409,'This email is already registered. Sign in or use your recovery code.')
    return {'user':{'id':uid,'email':email,'name':body.name.strip()},'token':store.session(uid),'recovery_code':recovery}


@app.post('/api/auth/login')
def login(body:Login,request:Request):
    limit('login:'+request.client.host,15,60)
    limit('account:'+str(body.email).lower(),15,300)
    with store.db() as c:
        user=c.execute('SELECT * FROM users WHERE email=?',(str(body.email).lower(),)).fetchone()
    if not user or not store.verify(body.password,user['password']):
        if not user:
            store.password_hash(body.password)
        raise HTTPException(401,'The email or password is incorrect.')
    return {'token':store.session(user['id']),'user':{x:user[x] for x in ['id','email','name']}}


@app.post('/api/auth/recover')
def recover(body:Recovery,request:Request):
    limit('recover:'+request.client.host,5,600)
    import hmac
    with store.db() as c:
        user=c.execute('SELECT * FROM users WHERE email=?',(str(body.email).lower(),)).fetchone()
        actual=hashlib.sha256(body.recovery_code.encode()).hexdigest()
        if not user or not hmac.compare_digest(actual,user['recovery_hash']):
            raise HTTPException(401,'The email or recovery code is incorrect.')
        code=store.new_id()+store.new_id()[:12]
        c.execute('UPDATE users SET password=?,recovery_hash=? WHERE id=?',(store.password_hash(body.password),hashlib.sha256(code.encode()).hexdigest(),user['id']))
        c.execute('DELETE FROM sessions WHERE user_id=?',(user['id'],))
    return {'recovery_code':code,'message':'Password reset. Save the new recovery code, then sign in.'}


@app.get('/api/auth/me')
def me(user=Depends(current_user)): return user


@app.post('/api/auth/logout')
def logout(credentials:HTTPAuthorizationCredentials=Depends(auth),user=Depends(current_user)):
    with store.db() as c:
        c.execute('DELETE FROM sessions WHERE token_hash=?',(hashlib.sha256(credentials.credentials.encode()).hexdigest(),))
    return {'ok':True}


@app.get('/api/health')
def health():
    return {'status':'ok','app':'niyam','version':'1.0.0','daily_only':True,'connections':connection_status(),'local_backend':True,'daily_feed':feed.status()}


@app.get('/api/data')
def data_status(user=Depends(current_user)): return catalog()


@app.get('/api/templates')
def templates(): return {k:template(k).model_dump(mode='json') for k in ['trend','rsi','breakout']}


@app.get('/api/references')
def references(): return REFERENCES


@app.post('/api/validate')
def validate(spec:Strategy,user=Depends(current_user)):
    return {'entry':describe(spec.entry),'exit':describe(spec.exit),'valid':True,
            'assumptions':['Daily long-only, next-open execution','Equal initial capital per selected stock','Nifty 50 price index benchmark'],
            'data':catalog()}


class Prompt(BaseModel):
    text:str=Field(min_length=3,max_length=12000)
    attachment_ids:list[str]=Field(default_factory=list,max_length=3)
    defaults:dict=Field(default_factory=dict)
    answers:list[dict]=Field(default_factory=list,max_length=8)


@app.post('/api/interpret')
def parse_strategy(body:Prompt,user=Depends(current_user)):
    limit('ai:'+user['id'],10,300)
    docs=[]
    with store.db() as c:
        for aid in body.attachment_ids:
            row=c.execute('SELECT name,body FROM attachments WHERE id=? AND user_id=?',(aid,user['id'])).fetchone()
            if not row:
                raise HTTPException(404,'An attached document was not found in your account.')
            docs.append(dict(row))
    if sum(len(x['body']) for x in docs)>60000:
        raise HTTPException(422,'Attach shorter documents or only the relevant pages (60,000 extracted characters total).')
    try:
        return interpret(body.text,docs,body.defaults,body.answers)
    except Exception:
        raise HTTPException(503,'The AI provider did not respond successfully. Check its model, endpoint, credits and key in .env, then retry. Your draft is still here.')


@app.post('/api/attachments')
async def upload(file:UploadFile=File(...),user=Depends(current_user)):
    limit('upload:'+user['id'],20,600)
    content=await file.read(8*1024*1024+1)
    await file.close()
    if len(content)>8*1024*1024:
        raise HTTPException(413,'Attach a document smaller than 8 MB.')
    name=Path(file.filename or 'document').name[:150]
    ext=Path(name).suffix.lower()
    try:
        if ext=='.pdf':
            from pypdf import PdfReader
            pdf=PdfReader(io.BytesIO(content))
            if pdf.is_encrypted or len(pdf.pages)>40:
                raise ValueError('Use an unlocked PDF of at most 40 pages.')
            text='\n'.join(f'[Page {i+1}] '+(p.extract_text() or '') for i,p in enumerate(pdf.pages))
        elif ext=='.docx':
            from docx import Document
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(x.file_size for x in archive.infolist())>30*1024*1024:
                    raise ValueError('The expanded Word document is too large.')
            doc=Document(io.BytesIO(content))
            text='\n'.join([p.text for p in doc.paragraphs]+[' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows])
        elif ext in {'.txt','.md'}:
            text=content.decode('utf-8-sig')
        else:
            raise ValueError('Please attach PDF, DOCX, TXT or Markdown.')
        if len(text.strip())<30:
            raise ValueError('No readable strategy text was found. Use a text-based PDF or paste the rules; scanned-image OCR is not enabled.')
        if len(text)>60000:
            raise ValueError('This document is too long. Attach only the relevant pages (60,000 characters maximum).')
    except ValueError as error:
        raise HTTPException(422,str(error))
    except Exception:
        raise HTTPException(422,'This document could not be read. Try a text PDF, DOCX or plain text.')
    aid=store.new_id()
    with store.db() as c:
        c.execute('INSERT INTO attachments VALUES(?,?,?,?,?)',(aid,user['id'],name,text,time.time()))
    return {'id':aid,'name':name,'characters':len(text),'preview':text[:400]}


@app.delete('/api/attachments/{aid}')
def remove_attachment(aid:str,user=Depends(current_user)):
    with store.db() as c:
        c.execute('DELETE FROM attachments WHERE id=? AND user_id=?',(aid,user['id']))
    return {'ok':True}


@app.get('/api/strategies')
def list_strategies(user=Depends(current_user)):
    with store.db() as c:
        return [{**dict(row),'spec':json.loads(row['spec'])} for row in c.execute('SELECT id,name,spec,updated FROM strategies WHERE user_id=? ORDER BY updated DESC',(user['id'],))]


@app.post('/api/strategies')
def save_strategy(spec:Strategy,user=Depends(current_user)):
    sid=store.new_id()
    with store.db() as c:
        if c.execute('SELECT COUNT(*) FROM strategies WHERE user_id=?',(user['id'],)).fetchone()[0]>=200:
            raise HTTPException(422,'You have reached 200 saved strategies. Remove an older draft first.')
        c.execute('INSERT INTO strategies VALUES(?,?,?,?,?)',(sid,user['id'],spec.name,spec.model_dump_json(),time.time()))
    return {'id':sid}


@app.delete('/api/strategies/{sid}')
def delete_strategy(sid:str,user=Depends(current_user)):
    with store.db() as c:
        c.execute('DELETE FROM strategies WHERE id=? AND user_id=?',(sid,user['id']))
    return {'ok':True}


@app.post('/api/jobs',status_code=202)
def submit(body:RunRequest,user=Depends(current_user)):
    limit('job:'+user['id'],20,600)
    if not body.strategy.symbols:
        raise HTTPException(422,'Choose at least one stock before starting.')
    if not body.strategy.price_basis_confirmed:
        raise HTTPException(422,'Confirm the price adjustment basis in Data assumptions first.')
    status=catalog()
    if not status['ready']:
        raise HTTPException(422,status['errors'][0] if status['errors'] else 'Add your prices and Nifty 50 benchmark data before running a backtest.')
    with jobs.SUBMIT_LOCK:
        with store.db() as c:
            if c.execute("SELECT COUNT(*) FROM jobs WHERE user_id=? AND state IN ('queued','running')",(user['id'],)).fetchone()[0]>=2:
                raise HTTPException(429,'Two jobs are already active. Wait for one to finish or cancel it.')
            if c.execute("SELECT COUNT(*) FROM jobs WHERE state IN ('queued','running')").fetchone()[0]>=10:
                raise HTTPException(429,'The research queue is full. Please try again shortly.')
            jid=store.new_id()
            c.execute('INSERT INTO jobs(id,user_id,state,request,created) VALUES(?,?,?,?,?)',(jid,user['id'],'queued',body.model_dump_json(),time.time()))
        jobs.POOL.submit(jobs.execute,jid)
    return {'id':jid}


@app.get('/api/jobs')
def list_jobs(user=Depends(current_user)):
    with store.db() as c:
        return [{**dict(r),'request':json.loads(r['request'])} for r in c.execute('SELECT id,state,progress,request,error,created FROM jobs WHERE user_id=? ORDER BY created DESC LIMIT 100',(user['id'],))]


@app.get('/api/jobs/{jid}')
def get_job(jid:str,user=Depends(current_user)):
    with store.db() as c:
        row=c.execute('SELECT id,state,progress,request,result,error,created FROM jobs WHERE id=? AND user_id=?',(jid,user['id'])).fetchone()
    if not row:
        raise HTTPException(404,'This result was not found in your account.')
    result=dict(row)
    result['request']=json.loads(result['request'])
    result['result']=json.loads(result['result']) if result['result'] else None
    return result


@app.post('/api/jobs/{jid}/cancel')
def cancel_job(jid:str,user=Depends(current_user)):
    with store.db() as c:
        c.execute("UPDATE jobs SET cancel=1 WHERE id=? AND user_id=? AND state IN ('queued','running')",(jid,user['id']))
    return {'ok':True}


@app.get('/api/signals')
def signal_list(user=Depends(current_user)):
    with store.db() as c:
        watches=[{**dict(r),'spec':json.loads(r['spec'])} for r in c.execute('SELECT id,spec,active,last_date,error FROM watches WHERE user_id=? ORDER BY created DESC',(user['id'],))]
        alerts=[{**dict(r),'body':json.loads(r['body'])} for r in c.execute('SELECT id,bar_date,symbol,action,body,created FROM alerts WHERE user_id=? ORDER BY created DESC LIMIT 200',(user['id'],))]
    return {'watches':watches,'alerts':alerts}


@app.post('/api/signals')
def add_watch(spec:Strategy,user=Depends(current_user)):
    if not spec.symbols or not spec.price_basis_confirmed:
        raise HTTPException(422,'Choose your stocks and confirm price assumptions first.')
    frames,benchmark,_=load_dataset()
    latest=signals.latest_closed_date(benchmark)
    if not latest:
        raise HTTPException(422,'Import completed daily candles and Nifty 50 data first.')
    try:
        adjusted=Strategy.model_validate({**spec.model_dump(mode='json'),'end':str(latest)})
        run=__import__('app.engine',fromlist=['run_backtest']).run_backtest
        run(adjusted,frames,benchmark,finalise=False)
    except ValueError as e:
        raise HTTPException(422,str(e))
    wid=store.new_id()
    with store.db() as c:
        if c.execute('SELECT COUNT(*) FROM watches WHERE user_id=?',(user['id'],)).fetchone()[0]>=20:
            raise HTTPException(422,'Remove an older watch before adding another (20 per account).')
        c.execute('INSERT INTO watches(id,user_id,spec,active,last_date,created) VALUES(?,?,?,?,?,?)',(wid,user['id'],spec.model_dump_json(),1,str(latest),time.time()))
    return {'id':wid,'baseline':str(latest),'message':'Watch created. It will evaluate newly completed daily bars while the backend is running.'}


class Toggle(BaseModel): active:bool


@app.put('/api/signals/{wid}')
def toggle_watch(wid:str,body:Toggle,user=Depends(current_user)):
    with store.db() as c:
        c.execute('UPDATE watches SET active=? WHERE id=? AND user_id=?',(body.active,wid,user['id']))
    return {'ok':True}


@app.delete('/api/signals/{wid}')
def delete_watch(wid:str,user=Depends(current_user)):
    with store.db() as c:
        c.execute('DELETE FROM alerts WHERE watch_id=? AND user_id=?',(wid,user['id']))
        c.execute('DELETE FROM watches WHERE id=? AND user_id=?',(wid,user['id']))
    return {'ok':True}


@app.post('/api/signals/refresh')
def refresh_signals(user=Depends(current_user)):
    limit('signals:'+user['id'],3,60)
    signals.refresh()
    return signal_list(user)


@app.get('/config-template')
def config_template(): return FileResponse(ROOT/'.env.example',filename='.env.example',media_type='text/plain')


app.mount('/',StaticFiles(directory=ROOT/'web',html=True),name='app')
